"""Validación y lectura segura de documentos destinados a Gemini."""

from pathlib import Path
from zipfile import BadZipFile, ZipFile

from django.conf import settings
from pptx import Presentation
from pypdf import PdfReader


class DocumentValidationError(ValueError):
    pass


def validate_document(uploaded_file):
    if not uploaded_file:
        raise DocumentValidationError("Selecciona un archivo PDF o PPTX.")
    if uploaded_file.size <= 0:
        raise DocumentValidationError("El archivo está vacío.")
    if uploaded_file.size > settings.MAX_DOCUMENT_UPLOAD_SIZE:
        raise DocumentValidationError("El documento debe pesar 50 MB o menos.")

    extension = Path(uploaded_file.name).suffix.lower()
    if extension == ".pdf":
        _validate_pdf(uploaded_file)
        document_type = "pdf"
    elif extension == ".pptx":
        _validate_pptx(uploaded_file)
        document_type = "pptx"
    else:
        raise DocumentValidationError("Solo se admiten archivos PDF y PPTX.")
    uploaded_file.seek(0)
    return document_type


def _validate_pdf(uploaded_file):
    header = uploaded_file.read(5)
    uploaded_file.seek(0)
    if header != b"%PDF-":
        raise DocumentValidationError("El archivo no es un PDF válido.")
    try:
        reader = PdfReader(uploaded_file, strict=False)
        if reader.is_encrypted:
            raise DocumentValidationError("No se admiten PDF protegidos con contraseña.")
        page_count = len(reader.pages)
    except DocumentValidationError:
        raise
    except Exception as error:
        raise DocumentValidationError("El archivo no es un PDF legible.") from error
    finally:
        uploaded_file.seek(0)
    if not 1 <= page_count <= 1000:
        raise DocumentValidationError("El PDF debe contener entre 1 y 1000 páginas.")


def _validate_pptx(uploaded_file):
    try:
        with ZipFile(uploaded_file) as archive:
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names or "ppt/presentation.xml" not in names:
                raise DocumentValidationError("El archivo no es un PPTX válido.")
        uploaded_file.seek(0)
        presentation = Presentation(uploaded_file)
        slide_count = len(presentation.slides)
    except DocumentValidationError:
        raise
    except (BadZipFile, KeyError, ValueError, OSError) as error:
        raise DocumentValidationError("El archivo no es un PPTX legible.") from error
    finally:
        uploaded_file.seek(0)
    if not 1 <= slide_count <= 200:
        raise DocumentValidationError("El PPTX debe contener entre 1 y 200 diapositivas.")


def extract_pptx_text(file_object):
    """Extrae texto, tablas y notas. El archivo nunca se publica ni se envía como PPTX."""
    file_object.seek(0)
    presentation = Presentation(file_object)
    sections = []
    for index, slide in enumerate(presentation.slides, start=1):
        fragments = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "").strip()
            if text:
                fragments.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        fragments.append(" | ".join(cells))
        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        except (AttributeError, ValueError):
            notes = ""
        if notes:
            fragments.append(f"Notas: {notes}")
        sections.append(f"DIAPOSITIVA {index}\n" + "\n".join(fragments))
    content = "\n\n".join(sections).strip()
    if not content:
        raise DocumentValidationError("El PPTX no contiene texto que se pueda convertir.")
    if len(content) > 150_000:
        raise DocumentValidationError("El PPTX contiene demasiado texto para una sola conversión.")
    return content
