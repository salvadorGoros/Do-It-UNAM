import json
import re
import tempfile
import wave
from io import BytesIO
from zipfile import ZipFile
from copy import deepcopy
from datetime import timedelta
from types import SimpleNamespace
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core import mail
from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from PIL import Image
from pptx import Presentation
from reportlab.pdfgen import canvas

from .exports import _evaluate_object, exportar_imagenes, exportar_video
from .gemini import GeminiResult, GeminiServiceError, TokenUsage
from .models import (
    DocumentoFuente,
    PerfilUsuario,
    PlantillaAnimacion,
    Proyecto,
    Rol,
    TrabajoConversion,
    UsoIA,
)
from .tasks import convertir_documento
from .scene import imagenes_disponibles, validar_diapositivas


class ProyectoTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.usuario = get_user_model().objects.create_user("ana", password="Prueba-local-123!")
        cls.otro = get_user_model().objects.create_user("luis", password="Prueba-local-456!")
        cls.proyecto = Proyecto.objects.create(nombre="Proyecto de Ana", propietario=cls.usuario)
        cls.ajeno = Proyecto.objects.create(nombre="Privado de Luis", propietario=cls.otro)

    def setUp(self):
        self.client.force_login(self.usuario)
        self.detalle = reverse("editor:detalle", args=[self.proyecto.pk])
        self.guardar = reverse("editor:guardar", args=[self.proyecto.pk])
        self.crear = reverse("editor:crear")
        self.subir_audio = reverse("editor:subir_audio", args=[self.proyecto.pk])
        self.estado_ia = reverse("editor:estado_ia", args=[self.proyecto.pk])
        self.generar_ia = reverse("editor:generar_ia", args=[self.proyecto.pk])
        self.subir_documento_ia = reverse("editor:subir_documento_ia", args=[self.proyecto.pk])

    def test_rutas_requieren_sesion(self):
        self.client.logout()
        for url, method in [
            ("/", "get"), (self.detalle, "get"), (self.crear, "post"),
            (self.guardar, "post"), (self.subir_audio, "post"),
            (reverse("editor:audios_subidos"), "get"),
            (self.estado_ia, "get"), (self.generar_ia, "post"),
            (self.subir_documento_ia, "post"),
        ]:
            with self.subTest(url=url):
                response = getattr(self.client, method)(url)
                self.assertRedirects(response, reverse("login") + "?next=" + url)

    def test_login_y_logout_solo_post(self):
        self.client.logout()
        response = self.client.post(reverse("login"), {"username": "ana", "password": "incorrecta"})
        self.assertTrue(response.context["form"].errors)
        response = self.client.post(reverse("login"), {
            "username": "ana", "password": "Prueba-local-123!", "next": self.detalle,
        })
        self.assertRedirects(response, self.detalle)
        self.assertEqual(self.client.get(reverse("logout")).status_code, 405)
        self.assertIn("_auth_user_id", self.client.session)
        self.assertRedirects(self.client.post(reverse("logout")), reverse("login"))
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_login_no_redirige_a_sitio_externo(self):
        self.client.logout()
        response = self.client.post(reverse("login"), {
            "username": "ana", "password": "Prueba-local-123!", "next": "https://example.com/",
        })
        self.assertRedirects(response, "/")

    def test_registro_crea_usuario_inicia_sesion_y_separa_proyectos(self):
        self.client.logout()
        url = reverse("editor:registro")
        self.assertEqual(self.client.get(url).status_code, 200)
        response = self.client.post(url, {
            "correo_institucional": "mariana@comunidad.unam.mx",
            "numero_cuenta": "312345678",
            "password1": "Registro-Segura-456!",
            "password2": "Registro-Segura-456!",
            "rol": "administrador",
        })
        self.assertRedirects(response, reverse("editor:index"))

        usuario = get_user_model().objects.get(username="mariana@comunidad.unam.mx")
        self.assertTrue(usuario.check_password("Registro-Segura-456!"))
        self.assertEqual(usuario.email, "mariana@comunidad.unam.mx")
        self.assertEqual(usuario.perfil_institucional.numero_cuenta, "312345678")
        self.assertEqual(usuario.perfil_institucional.rol.clave, Rol.Clave.ALUMNO)
        self.assertIn("_auth_user_id", self.client.session)

        listado = self.client.get(reverse("editor:index"))
        self.assertNotContains(listado, self.proyecto.nombre)
        self.assertNotContains(listado, self.ajeno.nombre)
        creado = self.client.post(reverse("editor:crear"), {"nombre": "Proyecto de Mariana"})
        proyecto = Proyecto.objects.get(nombre="Proyecto de Mariana")
        self.assertEqual(proyecto.propietario, usuario)
        self.assertRedirects(creado, reverse("editor:detalle", args=[proyecto.pk]))

    def test_registro_rechaza_correo_repetido_no_institucional_y_contrasenas_distintas(self):
        self.client.logout()
        url = reverse("editor:registro")
        datos = {
            "correo_institucional": "registro@comunidad.unam.mx",
            "numero_cuenta": "",
            "password1": "Registro-Segura-456!",
            "password2": "Registro-Segura-456!",
        }
        self.assertEqual(self.client.post(url, datos).status_code, 302)
        self.client.logout()
        response = self.client.post(url, datos)
        self.assertEqual(response.status_code, 200)
        self.assertIn("correo_institucional", response.context["form"].errors)
        self.assertEqual(get_user_model().objects.filter(email="registro@comunidad.unam.mx").count(), 1)

        datos["correo_institucional"] = "externa@gmail.com"
        response = self.client.post(url, datos)
        self.assertEqual(response.status_code, 200)
        self.assertIn("correo_institucional", response.context["form"].errors)

        datos["correo_institucional"] = "nueva@comunidad.unam.mx"
        datos["password2"] = "Otra-Segura-456!"
        response = self.client.post(url, datos)
        self.assertEqual(response.status_code, 200)
        self.assertIn("password2", response.context["form"].errors)
        self.assertFalse(get_user_model().objects.filter(email="nueva@comunidad.unam.mx").exists())

    def test_inicio_institucional_acepta_correo_y_numero_y_puede_cerrar_transicion(self):
        self.client.logout()
        self.client.post(reverse("editor:registro"), {
            "correo_institucional": "acceso@comunidad.unam.mx",
            "numero_cuenta": "31-2345-679",
            "password1": "Registro-Segura-456!",
            "password2": "Registro-Segura-456!",
        })
        self.client.logout()
        login_url = reverse("login")
        response = self.client.post(login_url, {
            "username": "ACCESO@COMUNIDAD.UNAM.MX",
            "password": "Registro-Segura-456!",
        })
        self.assertRedirects(response, reverse("editor:index"))
        self.client.logout()
        response = self.client.post(login_url, {
            "username": "312345679",
            "password": "Registro-Segura-456!",
        })
        self.assertRedirects(response, reverse("editor:index"))

        self.client.logout()
        with override_settings(ALLOW_LEGACY_USERNAME_LOGIN=False):
            response = self.client.post(login_url, {
                "username": "ana",
                "password": "Prueba-local-123!",
            })
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_numero_de_cuenta_es_unico_y_se_normaliza(self):
        self.client.logout()
        url = reverse("editor:registro")
        base = {
            "correo_institucional": "uno@comunidad.unam.mx",
            "numero_cuenta": "31-2345-680",
            "password1": "Registro-Segura-456!",
            "password2": "Registro-Segura-456!",
        }
        self.assertEqual(self.client.post(url, base).status_code, 302)
        profile = PerfilUsuario.objects.get(correo_institucional="uno@comunidad.unam.mx")
        self.assertEqual(profile.numero_cuenta, "312345680")
        self.client.logout()
        base["correo_institucional"] = "dos@comunidad.unam.mx"
        response = self.client.post(url, base)
        self.assertEqual(response.status_code, 200)
        self.assertIn("numero_cuenta", response.context["form"].errors)

    @override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
    def test_recuperacion_de_contrasena_por_correo_institucional(self):
        self.client.logout()
        self.client.post(reverse("editor:registro"), {
            "correo_institucional": "recuperar@comunidad.unam.mx",
            "numero_cuenta": "",
            "password1": "Registro-Segura-456!",
            "password2": "Registro-Segura-456!",
        })
        self.client.logout()
        reset_url = reverse("password_reset")
        response = self.client.post(reset_url, {"email": "recuperar@comunidad.unam.mx"})
        self.assertRedirects(response, reverse("password_reset_done"))
        self.assertEqual(len(mail.outbox), 1)
        reset_path = re.search(r"http://testserver(\S+)", mail.outbox[0].body).group(1)
        response = self.client.get(reset_path)
        self.assertEqual(response.status_code, 302)
        confirm_path = response.url
        response = self.client.post(confirm_path, {
            "new_password1": "Nueva-Segura-789!",
            "new_password2": "Nueva-Segura-789!",
        })
        self.assertRedirects(response, reverse("password_reset_complete"))
        response = self.client.post(reverse("login"), {
            "username": "recuperar@comunidad.unam.mx",
            "password": "Nueva-Segura-789!",
        })
        self.assertRedirects(response, reverse("editor:index"))

        self.client.logout()
        mail.outbox.clear()
        response = self.client.post(reset_url, {"email": "desconocido@comunidad.unam.mx"})
        self.assertRedirects(response, reverse("password_reset_done"))
        self.assertEqual(mail.outbox, [])

    def test_roles_base_existen_y_administrador_recibe_permisos_globales(self):
        self.assertEqual(
            list(Rol.objects.values_list("clave", "cuota_tokens_dia")),
            [("alumno", 10_000), ("docente", 50_000), ("administrador", 200_000)],
        )
        administrator = Rol.objects.get(clave=Rol.Clave.ADMINISTRADOR)
        profile = PerfilUsuario.objects.create(
            usuario=self.usuario,
            rol=administrator,
            correo_institucional="ana@comunidad.unam.mx",
            numero_cuenta="312345681",
        )
        self.usuario.refresh_from_db()
        self.assertTrue(self.usuario.is_staff)
        self.assertTrue(self.usuario.has_perm("auth.view_user"))
        self.assertTrue(self.usuario.has_perm("editor.change_rol"))
        self.assertContains(self.client.get(reverse("editor:index")), 'class="admin-access"')
        self.assertContains(self.client.get(self.detalle), 'class="admin-menu-link"')

        profile.rol = Rol.objects.get(clave=Rol.Clave.ALUMNO)
        profile.save()
        self.usuario.refresh_from_db()
        self.assertFalse(self.usuario.is_staff)
        self.assertFalse(self.usuario.has_perm("auth.view_user"))
        self.assertNotContains(self.client.get(reverse("editor:index")), 'class="admin-access"')
        self.assertNotContains(self.client.get(self.detalle), 'class="admin-menu-link"')

    def test_panel_administracion_crea_usuarios_cambia_roles_y_cuotas(self):
        administrator = Rol.objects.get(clave=Rol.Clave.ADMINISTRADOR)
        teacher = Rol.objects.get(clave=Rol.Clave.DOCENTE)
        student = Rol.objects.get(clave=Rol.Clave.ALUMNO)
        own_profile = PerfilUsuario.objects.create(
            usuario=self.usuario,
            rol=administrator,
            correo_institucional="ana@comunidad.unam.mx",
            numero_cuenta="312345681",
        )
        dashboard_url = reverse("editor:administracion")
        response = self.client.get(dashboard_url)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Usuarios y roles")

        response = self.client.post(dashboard_url, {
            "accion": "crear_usuario",
            "nuevo-correo_institucional": "docente@comunidad.unam.mx",
            "nuevo-numero_cuenta": "31-2345-682",
            "nuevo-password1": "Docente-Segura-456!",
            "nuevo-password2": "Docente-Segura-456!",
        })
        self.assertRedirects(response, dashboard_url)
        created = PerfilUsuario.objects.get(correo_institucional="docente@comunidad.unam.mx")
        self.assertEqual(created.rol, student)
        self.assertEqual(created.numero_cuenta, "312345682")

        response = self.client.post(dashboard_url, {
            "accion": "cambiar_rol",
            "perfil_id": created.pk,
            "rol_id": administrator.pk,
        })
        self.assertRedirects(response, dashboard_url)
        created.refresh_from_db()
        created.usuario.refresh_from_db()
        self.assertEqual(created.rol, administrator)
        self.assertTrue(created.usuario.is_staff)

        self.client.post(dashboard_url, {
            "accion": "actualizar_cuota",
            "rol_id": teacher.pk,
            "cuota_tokens_dia": "64000",
        })
        teacher.refresh_from_db()
        self.assertEqual(teacher.cuota_tokens_dia, 64_000)

        response = self.client.post(dashboard_url, {
            "accion": "cambiar_rol",
            "perfil_id": own_profile.pk,
            "rol_id": student.pk,
        }, follow=True)
        own_profile.refresh_from_db()
        self.assertEqual(own_profile.rol, administrator)
        self.assertContains(response, "No puedes retirar tu propio acceso administrativo")

        self.client.force_login(self.otro)
        self.assertEqual(self.client.get(dashboard_url).status_code, 403)

    @override_settings(GEMINI_API_KEY="prueba")
    def test_estado_ia_usa_cuota_configurable_del_rol(self):
        teacher = Rol.objects.get(clave=Rol.Clave.DOCENTE)
        teacher.cuota_tokens_dia = 43_210
        teacher.save(update_fields=["cuota_tokens_dia"])
        PerfilUsuario.objects.create(
            usuario=self.usuario,
            rol=teacher,
            correo_institucional="docente@comunidad.unam.mx",
        )
        response = self.client.get(self.estado_ia)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["role"], {"key": "docente", "name": "Docente"})
        self.assertEqual(response.json()["quota"]["limit"], 43_210)

    def test_listado_solo_propios_y_orden_modificacion(self):
        nuevo = Proyecto.objects.create(nombre="Otro de Ana", propietario=self.usuario)
        Proyecto.objects.filter(pk=self.proyecto.pk).update(modificado=timezone.now() - timedelta(days=1))
        response = self.client.get("/")
        self.assertQuerySetEqual(response.context["proyectos"], [nuevo, self.proyecto])
        self.assertNotContains(response, self.ajeno.nombre)
        self.client.post(self.guardar, {"nombre": "Ahora más reciente"})
        self.assertEqual(self.client.get("/").context["proyectos"][0].pk, self.proyecto.pk)

    def test_panel_de_proyectos_incluye_busqueda_resumen_y_vista_previa(self):
        self.proyecto.diapositivas = [
            {
                "id": "slide-1",
                "color_fondo": "#fef3c7",
                "objetos": [self.cuadro(texto="Portada")],
                "fondo": "",
                "duracion_frames": 300,
            },
            {
                "id": "slide-2", "color_fondo": "#ffffff", "objetos": [],
                "fondo": "", "duracion_frames": 300,
            },
        ]
        self.proyecto.save(update_fields=["diapositivas"])

        response = self.client.get(reverse("editor:index"))
        proyecto = next(item for item in response.context["proyectos"] if item.pk == self.proyecto.pk)

        self.assertEqual(proyecto.numero_diapositivas, 2)
        self.assertEqual(proyecto.total_objetos, 1)
        self.assertEqual(proyecto.preview_data["color_fondo"], "#fef3c7")
        self.assertContains(response, 'id="project-search"')
        self.assertContains(response, f'id="project-preview-{self.proyecto.pk}"')
        self.assertContains(response, "Portada")
        self.assertContains(response, "2 diapositivas")
        self.assertNotContains(response, self.ajeno.nombre)

    def test_estado_vacio(self):
        self.usuario.proyectos.all().delete()
        self.assertContains(self.client.get("/"), "Aún no tienes proyectos")

    def test_logo_enlace_y_favicon_en_editor_y_proyectos(self):
        logo = "/static/editor/images/logo.png?v=4"
        proyectos = self.client.get("/")
        editor = self.client.get(self.detalle)
        favicon = '<link rel="icon" type="image/png" href="/static/editor/images/favicon.png?v=2">'
        self.assertContains(proyectos, favicon)
        self.assertContains(proyectos, f'<img class="brand-logo" src="{logo}" alt="DO IT UNAM" width="52" height="52">')
        self.assertContains(editor, favicon)
        self.assertContains(editor, f'<img src="{logo}" alt="DO IT UNAM" width="52" height="52">')
        self.assertContains(editor, 'href="/" class="projects-logo" aria-label="Regresar a Mis proyectos"')

    def test_crear_ignora_propietario_enviado_y_permite_repetidos(self):
        for _ in range(2):
            response = self.client.post(self.crear, {"nombre": "  Diseño nuevo  ", "propietario": self.otro.pk})
            proyecto = Proyecto.objects.filter(propietario=self.usuario).first()
            self.assertEqual(proyecto.nombre, "Diseño nuevo")
            self.assertRedirects(response, reverse("editor:detalle", args=[proyecto.pk]))
        self.assertEqual(self.usuario.proyectos.filter(nombre="Diseño nuevo").count(), 2)

    def test_guardar_persiste_y_muestra_confirmacion(self):
        creado = self.proyecto.creado
        response = self.client.post(self.guardar, {"nombre": "  Nuevo nombre  ", "propietario": self.otro.pk}, follow=True)
        self.assertContains(response, "Proyecto guardado")
        self.proyecto.refresh_from_db()
        self.assertEqual(self.proyecto.nombre, "Nuevo nombre")
        self.assertEqual(self.proyecto.propietario, self.usuario)
        self.assertEqual(self.proyecto.creado, creado)
        self.assertContains(self.client.get(self.detalle), 'value="Nuevo nombre"')

    def test_nombres_invalidos_no_escriben_y_conservan_entrada(self):
        for nombre in ["", "   ", "x" * 151]:
            for url in [self.crear, self.guardar]:
                with self.subTest(nombre=nombre, url=url):
                    cantidad = Proyecto.objects.count()
                    response = self.client.post(url, {"nombre": nombre})
                    self.assertEqual(response.status_code, 200)
                    self.assertIn("nombre", response.context["form"].errors)
                    self.assertEqual(response.context["form"]["nombre"].value(), nombre)
                    self.assertEqual(Proyecto.objects.count(), cantidad)
                    self.proyecto.refresh_from_db()
                    self.assertEqual(self.proyecto.nombre, "Proyecto de Ana")

    def test_limite_y_escape_html(self):
        self.assertEqual(self.client.post(self.guardar, {"nombre": "á" * 150}).status_code, 302)
        response = self.client.post(self.guardar, {"nombre": '<script>alert("x")</script>'}, follow=True)
        self.assertNotContains(response, '<script>alert("x")</script>')
        self.assertContains(response, "&lt;script&gt;")

    def test_ajenos_e_inexistentes_devuelven_404(self):
        for pk in [self.ajeno.pk, 99999]:
            with self.subTest(pk=pk):
                self.assertEqual(self.client.get(reverse("editor:detalle", args=[pk])).status_code, 404)
                self.assertEqual(self.client.post(reverse("editor:guardar", args=[pk]), {"nombre": "Intrusión"}).status_code, 404)
        self.ajeno.refresh_from_db()
        self.assertEqual(self.ajeno.nombre, "Privado de Luis")

    def test_get_no_modifica_proyectos(self):
        estado = list(Proyecto.objects.values())
        for url in [self.crear, self.guardar]:
            self.assertEqual(self.client.get(url, {"nombre": "Cambio"}).status_code, 405)
        self.client.get("/")
        self.client.get(self.detalle)
        self.assertEqual(list(Proyecto.objects.values()), estado)

    def test_csrf_en_creacion_guardado_logout_y_login(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.usuario)
        for url in [
            self.crear, self.guardar, self.subir_audio, self.generar_ia,
            reverse("logout"), reverse("editor:registro"), reverse("password_reset"),
        ]:
            self.assertEqual(client.post(url, {"nombre": "Sin token"}).status_code, 403)
        client.get(self.detalle)
        token = client.cookies["csrftoken"].value
        response = client.post(self.guardar, {"nombre": "Con token", "csrfmiddlewaretoken": token})
        self.assertEqual(response.status_code, 302)
        client.logout()
        self.assertEqual(client.post(reverse("login"), {"username": "ana", "password": "Prueba-local-123!"}).status_code, 403)

    def cuadro(self, **changes):
        obj = {"id": "cuadro-1", "texto": "Hola\nmundo", "x": 50, "y": 60,
                "ancho": 300, "alto": 150, "fuente": "Poppins", "tamano": 36,
                "color": "#ffffff", "negrita": False, "cursiva": False,
                "subrayado": False, "alineacion": "center", **changes}
        obj.setdefault("fotogramas", [dict(frame=0, x=obj["x"], y=obj["y"], giro=0, opacidad=1, interpolacion="lineal")])
        return obj

    def test_proyecto_nuevo_sin_textos_fijos(self):
        self.assertEqual(self.proyecto.objetos, [])
        response = self.client.get(self.detalle)
        self.assertContains(response, 'id="text-tool"')
        self.assertNotContains(response, 'id="id_titulo"')

    def test_galerias_muestran_imagenes_disponibles(self):
        response = self.client.get(self.detalle)
        elementos = imagenes_disponibles("elementos")
        fondos = imagenes_disponibles("fondos")
        self.assertTrue(elementos)
        self.assertTrue(fondos)
        self.assertEqual(sum(len(grupo["assets"]) for grupo in response.context["elementos"]), len(elementos))
        self.assertEqual(sum(len(grupo["assets"]) for grupo in response.context["fondos"]), len(fondos))
        self.assertContains(response, f'data-asset-path="elementos/{elementos[0]}"')
        self.assertContains(response, f'data-asset-path="fondos/{fondos[0]}"')
        self.assertIn("Personajes de la FESC", [grupo["categoria"] for grupo in response.context["elementos"]])

    def test_subir_imagen_persiste_en_biblioteca_privada(self):
        image_data = BytesIO()
        Image.new("RGB", (12, 8), "red").save(image_data, format="PNG")
        image_data.seek(0)
        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            response = self.client.post(
                reverse("editor:subir_imagen", args=[self.proyecto.pk]),
                {"images": SimpleUploadedFile("mifoto.png", image_data.getvalue(), content_type="image/png")},
            )
            self.assertEqual(response.status_code, 201)
            uploaded = response.json()["images"][0]
            self.assertRegex(
                uploaded["path"],
                rf"^uploads/{self.usuario.pk}/[a-f0-9]{{32}}(?:_[a-zA-Z0-9_-]{{1,60}})?\.png$",
            )
            library = self.client.get(reverse("editor:imagenes_subidas"))
            self.assertEqual(library["Cache-Control"], "private, no-store")
            self.assertEqual(library.json()["images"], [uploaded])
            image_response = self.client.get(uploaded["url"])
            self.assertEqual(image_response.status_code, 200)
            self.assertEqual(image_response["Content-Type"], "image/png")
            self.assertEqual(b"".join(image_response.streaming_content), image_data.getvalue())

            objeto = self.cuadro(tipo="imagen", imagen=uploaded["path"], texto="")
            exported = exportar_imagenes([{
                "objetos": [objeto],
                "fondo": "",
                "color_fondo": "#ffffff",
            }])
            with ZipFile(BytesIO(exported)) as archive:
                with Image.open(BytesIO(archive.read("diapositiva-1.png"))) as slide_image:
                    self.assertEqual(slide_image.size, (1600, 900))

            otro_cliente = Client()
            otro_cliente.force_login(self.otro)
            self.assertEqual(otro_cliente.get(uploaded["url"]).status_code, 404)
            self.client.logout()
            self.assertEqual(self.client.get(uploaded["url"]).status_code, 302)

    def test_subida_rechaza_archivos_que_no_sean_imagenes(self):
        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            response = self.client.post(
                reverse("editor:subir_imagen", args=[self.proyecto.pk]),
                {"images": SimpleUploadedFile("imagen.png", b"esto no es una imagen", content_type="image/png")},
            )
            self.assertEqual(response.status_code, 400)
            self.assertIn("imagen válida", response.json()["error"])
            self.assertEqual(self.client.get(reverse("editor:imagenes_subidas")).json()["images"], [])

    def test_audio_se_sube_se_asigna_y_permanece_privado(self):
        audio_data = (
            b"RIFF" + (36).to_bytes(4, "little") + b"WAVE"
            + b"fmt " + (16).to_bytes(4, "little")
            + (1).to_bytes(2, "little") + (1).to_bytes(2, "little")
            + (8000).to_bytes(4, "little") + (8000).to_bytes(4, "little")
            + (1).to_bytes(2, "little") + (8).to_bytes(2, "little")
            + b"data" + (0).to_bytes(4, "little")
        )
        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            response = self.client.post(
                reverse("editor:subir_audio", args=[self.proyecto.pk]),
                {"audios": SimpleUploadedFile("narracion.wav", audio_data, content_type="audio/wav")},
            )
            self.assertEqual(response.status_code, 201)
            uploaded = response.json()["audios"][0]
            self.assertRegex(
                uploaded["path"],
                rf"^audio_uploads/{self.usuario.pk}/[a-f0-9]{{32}}(?:_[a-zA-Z0-9_-]{{1,60}})?\.wav$",
            )

            library = self.client.get(reverse("editor:audios_subidos"))
            self.assertEqual(library["Cache-Control"], "private, no-store")
            self.assertEqual(library.json()["audios"], [uploaded])
            audio_response = self.client.get(uploaded["url"])
            self.assertEqual(audio_response.status_code, 200)
            self.assertEqual(audio_response["Content-Type"], "audio/wav")
            self.assertEqual(b"".join(audio_response.streaming_content), audio_data)

            slide = {
                "id": "slide-1", "objetos": [], "fondo": "", "audio": uploaded["path"],
                "color_fondo": "#ffffff", "duracion_frames": 300,
            }
            save_response = self.client.post(self.guardar, {
                "nombre": "Con narración", "objetos": "[]", "fondo": "",
                "duracion_frames": 300, "diapositivas": json.dumps([slide]),
                "diapositiva_activa": 0,
            })
            self.assertRedirects(save_response, self.detalle)
            self.proyecto.refresh_from_db()
            self.assertEqual(self.proyecto.diapositivas[0]["audio"], uploaded["path"])

            otro_cliente = Client()
            otro_cliente.force_login(self.otro)
            self.assertEqual(otro_cliente.get(uploaded["url"]).status_code, 404)

    def test_subida_de_audio_rechaza_archivo_invalido(self):
        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            response = self.client.post(
                reverse("editor:subir_audio", args=[self.proyecto.pk]),
                {"audios": SimpleUploadedFile("engaño.mp3", b"no es audio", content_type="audio/mpeg")},
            )
            self.assertEqual(response.status_code, 400)
            self.assertIn("MP3, WAV, OGG y M4A", response.json()["error"])
            self.assertEqual(self.client.get(reverse("editor:audios_subidos")).json()["audios"], [])

    def test_guardado_rechaza_audio_subido_de_otro_usuario(self):
        audio_data = b"OggS" + b"\x00" * 64
        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            otro_cliente = Client()
            otro_cliente.force_login(self.otro)
            upload = otro_cliente.post(
                reverse("editor:subir_audio", args=[self.ajeno.pk]),
                {"audios": SimpleUploadedFile("ajeno.ogg", audio_data, content_type="audio/ogg")},
            )
            audio_path = upload.json()["audios"][0]["path"]
            slide = {
                "id": "slide-1", "objetos": [], "fondo": "", "audio": audio_path,
                "duracion_frames": 300,
            }
            response = self.client.post(self.guardar, {
                "nombre": "Audio ajeno", "objetos": "[]", "fondo": "",
                "duracion_frames": 300, "diapositivas": json.dumps([slide]),
                "diapositiva_activa": 0,
            })
            self.assertEqual(response.status_code, 200)
            self.assertIn("no pertenece", response.context["form"].errors["diapositivas"][0])
            self.proyecto.refresh_from_db()
            self.assertEqual(self.proyecto.diapositivas, [])

    def test_guardado_rechaza_imagen_subida_de_otro_usuario(self):
        image_data = BytesIO()
        Image.new("RGB", (4, 4), "blue").save(image_data, format="PNG")
        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            otro_cliente = Client()
            otro_cliente.force_login(self.otro)
            upload = otro_cliente.post(
                reverse("editor:subir_imagen", args=[self.ajeno.pk]),
                {"images": SimpleUploadedFile("ajena.png", image_data.getvalue(), content_type="image/png")},
            )
            image_path = upload.json()["images"][0]["path"]
            objeto = self.cuadro(tipo="imagen", imagen=image_path, texto="")
            response = self.client.post(self.guardar, {"nombre": "Imagen ajena", "objetos": json.dumps([objeto])})
            self.assertEqual(response.status_code, 200)
            self.assertIn("no pertenece", response.context["form"].errors["diapositivas"][0])
            self.proyecto.refresh_from_db()
            self.assertEqual(self.proyecto.objetos, [])

    def test_elementos_y_fondos_se_guardan_y_recuperan(self):
        elemento = self.cuadro(tipo="imagen", imagen=f"elementos/{imagenes_disponibles('elementos')[0]}", texto="")
        fondo = f"fondos/{imagenes_disponibles('fondos')[0]}"
        response = self.client.post(self.guardar, {
            "nombre": "Escena con imágenes",
            "objetos": json.dumps([elemento]),
            "fondo": fondo,
        })
        self.assertRedirects(response, self.detalle)
        self.proyecto.refresh_from_db()
        self.assertEqual(self.proyecto.objetos, [elemento])
        self.assertEqual(self.proyecto.fondo, fondo)
        editor = self.client.get(self.detalle)
        self.assertEqual(editor.context["form"]["fondo"].value(), fondo)
        self.assertContains(editor, f'data-asset-path="{elemento["imagen"]}"')

    def test_diapositivas_guardan_contenido_independiente_y_duracion(self):
        primero = self.cuadro(texto="Primera hoja")
        segundo = self.cuadro(id="cuadro-2", texto="Segunda hoja", x=300)
        slides = [
            {
                "id": "slide-1", "objetos": [primero], "fondo": "", "color_fondo": "#dbeafe",
                "duracion_frames": 300, "notas": "Guion editable del expositor",
            },
            {"id": "slide-2", "objetos": [segundo], "fondo": "", "duracion_frames": 600},
        ]
        response = self.client.post(self.guardar, {
            "nombre": "Presentación",
            "objetos": json.dumps([primero]),
            "fondo": "",
            "duracion_frames": 300,
            "diapositivas": json.dumps(slides),
            "diapositiva_activa": 0,
        })
        self.assertRedirects(response, self.detalle)
        self.proyecto.refresh_from_db()
        self.assertEqual(self.proyecto.diapositivas, slides)
        self.assertEqual(self.proyecto.objetos, [primero])

    def test_color_de_diapositiva_es_opcional_compatible_y_validado(self):
        slide = {"id": "slide-1", "objetos": [], "fondo": "", "duracion_frames": 300}
        validar_diapositivas([slide])
        validar_diapositivas([{**slide, "color_fondo": "#dbeafe"}])
        validar_diapositivas([{**slide, "notas": "Guion de la diapositiva"}])
        with self.assertRaises(ValidationError):
            validar_diapositivas([{**slide, "notas": "x" * 20_001}])
        for color in ("blue", "#fff", 123):
            with self.subTest(color=color), self.assertRaises(ValidationError):
                validar_diapositivas([{**slide, "color_fondo": color}])

    def test_audio_de_diapositiva_es_opcional_y_validado(self):
        slide = {"id": "slide-1", "objetos": [], "fondo": "", "duracion_frames": 300}
        validar_diapositivas([slide])
        validar_diapositivas([{
            **slide,
            "audio": "audio_uploads/1/0123456789abcdef0123456789abcdef_narracion.mp3",
        }])
        for audio in ("audio.mp3", "../privado.mp3", 123):
            with self.subTest(audio=audio), self.assertRaises(ValidationError):
                validar_diapositivas([{**slide, "audio": audio}])

    def test_exporta_todas_las_diapositivas_en_pdf_pptx_y_png(self):
        fondo = f"fondos/{imagenes_disponibles('fondos')[0]}"
        objeto = self.cuadro(texto="Editable en PowerPoint")
        objeto["fotogramas"][0]["giro"] = 90
        self.proyecto.diapositivas = [
            {"id": "slide-1", "objetos": [objeto], "fondo": fondo, "color_fondo": "#dbeafe", "duracion_frames": 300},
            {"id": "slide-2", "objetos": [], "fondo": "", "color_fondo": "#fef9c3", "duracion_frames": 600},
        ]
        self.proyecto.save()

        pdf = self.client.get(reverse("editor:exportar", args=[self.proyecto.pk, "pdf"]))
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        pptx = self.client.get(reverse("editor:exportar", args=[self.proyecto.pk, "pptx"]))
        deck = Presentation(BytesIO(pptx.content))
        self.assertEqual(len(deck.slides), 2)
        self.assertIn("Editable en PowerPoint", deck.slides[0].shapes[1].text)
        self.assertEqual(deck.slides[0].shapes[1].rotation, 90)
        self.assertEqual(tuple(deck.slides[0].background.fill.fore_color.rgb), (219, 234, 254))
        self.assertEqual(tuple(deck.slides[1].background.fill.fore_color.rgb), (254, 249, 195))
        images = self.client.get(reverse("editor:exportar", args=[self.proyecto.pk, "png"]))
        with ZipFile(BytesIO(images.content)) as archive:
            self.assertEqual(archive.namelist(), ["diapositiva-1.png", "diapositiva-2.png"])
            rotated_png = archive.read("diapositiva-1.png")
            with Image.open(BytesIO(archive.read("diapositiva-2.png"))) as slide_image:
                self.assertEqual(slide_image.getpixel((0, 0)), (254, 249, 195))
        object_without_rotation = deepcopy(objeto)
        object_without_rotation["fotogramas"][0]["giro"] = 0
        unrotated_slides = [dict(self.proyecto.diapositivas[0], objetos=[object_without_rotation])]
        with ZipFile(BytesIO(exportar_imagenes(unrotated_slides))) as archive:
            unrotated_png = archive.read("diapositiva-1.png")
        self.assertNotEqual(rotated_png, unrotated_png)

    def test_exportacion_requiere_propietario_y_formato_valido(self):
        self.client.force_login(self.otro)
        self.assertEqual(self.client.get(reverse("editor:exportar", args=[self.proyecto.pk, "pdf"])).status_code, 404)
        self.client.force_login(self.usuario)
        self.assertEqual(self.client.get(reverse("editor:exportar", args=[self.proyecto.pk, "exe"])).status_code, 404)

    def test_autoguardado_ajax_guarda_y_reporta_errores(self):
        objeto = self.cuadro(animacion="subir")
        response = self.client.post(
            self.guardar,
            {
                "nombre": "Guardado automático",
                "objetos": json.dumps([objeto]),
                "fondo": "",
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"saved": True})
        self.proyecto.refresh_from_db()
        self.assertEqual(self.proyecto.nombre, "Guardado automático")
        self.assertEqual(self.proyecto.objetos, [objeto])

        objeto["animacion"] = "efecto-inexistente"
        response = self.client.post(
            self.guardar,
            {"nombre": "No guardar", "objetos": json.dumps([objeto]), "fondo": ""},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("objetos", response.json()["errors"])
        self.proyecto.refresh_from_db()
        self.assertEqual(self.proyecto.nombre, "Guardado automático")

    def test_editor_muestra_preajustes_y_linea_de_tiempo_personalizable(self):
        response = self.client.get(self.detalle)
        for preset in ["ninguna", "aparecer", "subir", "zoom", "girar", "rebote"]:
            self.assertContains(response, f'data-animation="{preset}"')
        self.assertContains(response, 'id="animation-tool"')
        self.assertContains(response, 'id="preview-animation"')
        self.assertContains(response, 'id="timeline-panel"')
        self.assertContains(response, 'id="keyframe-easing"')
        self.assertContains(response, 'id="audio-waveform-tracks"')
        self.assertContains(response, 'id="save-animation-template"')
        self.assertNotContains(response, 'id="timeline-scroll"')
        self.assertNotContains(response, 'id="commit-key"')

    def test_rechaza_rutas_de_imagenes_que_no_pertenecen_a_las_galerias(self):
        objeto = self.cuadro(tipo="imagen", imagen="../logo.png", texto="")
        response = self.client.post(self.guardar, {
            "nombre": "No guardar",
            "objetos": json.dumps([objeto]),
            "fondo": "",
        })
        self.assertIn("objetos", response.context["form"].errors)
        fondo_invalido = self.client.post(self.guardar, {
            "nombre": "No guardar",
            "objetos": "[]",
            "fondo": "elementos/bola.png",
        })
        self.assertIn("fondo", fondo_invalido.context["form"].errors)
        self.proyecto.refresh_from_db()
        self.assertEqual(self.proyecto.objetos, [])
        self.assertEqual(self.proyecto.fondo, "")

    def test_formulario_antiguo_conserva_fondo(self):
        fondo = f"fondos/{imagenes_disponibles('fondos')[0]}"
        self.proyecto.fondo = fondo
        self.proyecto.save()
        response = self.client.post(self.guardar, {"nombre": "Solo nombre"})
        self.assertRedirects(response, self.detalle)
        self.proyecto.refresh_from_db()
        self.assertEqual(self.proyecto.fondo, fondo)

    def test_guardar_cuadros_independientes_y_recuperar(self):
        objetos = [self.cuadro(), self.cuadro(id="cuadro-2", texto="Segundo", x=400, tamano=24, negrita=True)]
        response = self.client.post(self.guardar, {"nombre": "Escena", "objetos": json.dumps(objetos)}, follow=True)
        self.assertContains(response, "Proyecto guardado")
        self.proyecto.refresh_from_db()
        self.assertEqual(self.proyecto.objetos, objetos)
        self.assertEqual(self.client.get(self.detalle).context["objetos"], objetos)
        self.assertEqual(self.ajeno.objetos, [])

    def test_guardar_eliminacion_y_texto_vacio(self):
        for objetos in [[self.cuadro(texto="")], []]:
            self.assertEqual(self.client.post(self.guardar, {"nombre": "Escena", "objetos": json.dumps(objetos)}).status_code, 302)
            self.proyecto.refresh_from_db()
            self.assertEqual(self.proyecto.objetos, objetos)

    def test_rechaza_escenas_invalidas_sin_modificar_proyecto(self):
        invalidos = [None, {}, "texto", [self.cuadro(x=-8001)], [self.cuadro(x=8001)],
            [self.cuadro(ancho=8001)], [self.cuadro(alto=4501)],
            [self.cuadro(tamano=121)], [self.cuadro(tamano=True)], [self.cuadro(x=float("nan"))],
            [self.cuadro(fuente="<script>")], [self.cuadro(color="red")],
            [self.cuadro(texto="a" * 5001)], [self.cuadro(alto=0)],
            [self.cuadro(negrita="true")], [self.cuadro(id="<script>")],
            [self.cuadro(), self.cuadro()], [self.cuadro(extra="propiedad")],
            [self.cuadro(id=f"id-{i}") for i in range(101)]]
        # null es una entrada inválida: para vaciar el lienzo se envía [].
        for objetos in invalidos:
            with self.subTest(objetos=str(objetos)[:80]):
                response = self.client.post(self.guardar, {"nombre": "No guardar", "objetos": json.dumps(objetos)})
                self.assertIn("objetos", response.context["form"].errors)
                self.proyecto.refresh_from_db()
                self.assertEqual(self.proyecto.nombre, "Proyecto de Ana")
                self.assertEqual(self.proyecto.objetos, [])

    def test_objetos_y_fotogramas_pueden_salir_del_marco(self):
        objeto = self.cuadro(x=-120, y=470, ancho=1000, alto=600)
        objeto["fotogramas"][0].update(x=-120, y=470)
        response = self.client.post(self.guardar, {
            "nombre": "Objeto fuera del marco",
            "objetos": json.dumps([objeto]),
        })
        self.assertEqual(response.status_code, 302)
        self.proyecto.refresh_from_db()
        self.assertEqual(self.proyecto.objetos, [objeto])

    def test_json_roto_y_error_nombre_conservan_cuadros(self):
        response = self.client.post(self.guardar, {"nombre": "Escena", "objetos": "{roto"})
        self.assertIn("objetos", response.context["form"].errors)
        objetos = [self.cuadro()]
        response = self.client.post(self.guardar, {"nombre": "", "objetos": json.dumps(objetos)})
        self.assertEqual(response.context["objetos"], objetos)
        self.proyecto.refresh_from_db()
        self.assertEqual(self.proyecto.objetos, [])

    def test_html_en_cuadros_se_serializa_sin_ejecutar(self):
        texto = '</script><img src=x onerror=alert(1)>'
        response = self.client.post(self.guardar, {"nombre": "Escena", "objetos": json.dumps([self.cuadro(texto=texto)])}, follow=True)
        self.assertNotContains(response, texto)
        self.proyecto.refresh_from_db()
        self.assertEqual(self.proyecto.objetos[0]["texto"], texto)

    def test_nombre_solo_no_borra_cuadros_y_ajeno_no_puede_guardar(self):
        self.proyecto.objetos = [self.cuadro()]
        self.proyecto.save()
        self.client.post(self.guardar, {"nombre": "Solo nombre"})
        self.proyecto.refresh_from_db()
        self.assertEqual(self.proyecto.objetos, [self.cuadro()])

    def test_animacion_y_duracion_persisten(self):
        objeto = self.cuadro()
        objeto["fotogramas"].append(dict(frame=600, x=300, y=200, giro=-180, opacidad=0.25, interpolacion="suave"))
        response = self.client.post(self.guardar, {"nombre": "Animación", "objetos": json.dumps([objeto]), "duracion_frames": 900})
        self.assertEqual(response.status_code, 302)
        self.proyecto.refresh_from_db()
        self.assertEqual(self.proyecto.duracion_frames, 900)
        self.assertEqual(self.proyecto.objetos, [objeto])

    def test_rechaza_fotogramas_invalidos(self):
        base = self.cuadro()
        key = base["fotogramas"][0]
        invalidos = [[], [key, key], [{**key, "frame": -1}], [{**key, "frame": 1.5}],
            [{**key, "frame": True}], [{**key, "frame": 108001}], [{**key, "x": 8001}],
            [{**key, "giro": 361}], [{**key, "opacidad": -0.1}], [{**key, "opacidad": float("nan")}],
            [{**key, "opacidad": True}], [{**key, "interpolacion": "desconocida"}],
            [{**key, "frame": 10}, key], [{**key, "extra": 3}]]
        for keys in invalidos:
            with self.subTest(keys=keys):
                response = self.client.post(self.guardar, {"nombre": "Inválido", "objetos": json.dumps([{**base, "fotogramas": keys}])})
                self.assertTrue(response.context["form"].errors)
                self.proyecto.refresh_from_db()
                self.assertEqual(self.proyecto.objetos, [])

    def test_duracion_invalida_no_guarda_ningun_cambio(self):
        objeto = self.cuadro()
        objeto["fotogramas"].append({**objeto["fotogramas"][0], "frame": 300})
        for duracion in [29, 108001, 60, "abc", 30.5]:
            response = self.client.post(self.guardar, {"nombre": "Inválido", "objetos": json.dumps([objeto]), "duracion_frames": duracion})
            self.assertIn("duracion_frames", response.context["form"].errors)
            self.proyecto.refresh_from_db()
            self.assertEqual(self.proyecto.nombre, "Proyecto de Ana")

    def test_formulario_antiguo_conserva_pistas_por_id(self):
        objeto = self.cuadro()
        objeto["fotogramas"].append(dict(frame=100, x=200, y=100, giro=90, opacidad=0.5, interpolacion="lineal"))
        self.proyecto.objetos = [objeto]
        self.proyecto.duracion_frames = 600
        self.proyecto.save()
        anterior = deepcopy(objeto)
        del anterior["fotogramas"]
        anterior["texto"] = "Edición desde pestaña antigua"
        response = self.client.post(self.guardar, {"nombre": "Antiguo", "objetos": json.dumps([anterior])})
        self.assertEqual(response.status_code, 302)
        self.proyecto.refresh_from_db()
        self.assertEqual(self.proyecto.objetos[0]["fotogramas"], objeto["fotogramas"])
        self.assertEqual(self.proyecto.duracion_frames, 600)

    def test_ajeno_no_puede_cambiar_duracion_ni_pistas(self):
        self.client.force_login(self.otro)
        response = self.client.post(self.guardar, {"nombre": "Ajeno", "duracion_frames": 900, "objetos": json.dumps([self.cuadro()])})
        self.assertEqual(response.status_code, 404)
        self.proyecto.refresh_from_db()
        self.assertEqual(self.proyecto.objetos, [])
        self.assertEqual(self.proyecto.duracion_frames, 300)

    def test_fotogramas_avanzados_audio_multipista_y_marcadores_son_validos(self):
        objeto = self.cuadro(animacion="pulso", fuente="Georgia")
        objeto["fotogramas"] = [
            {
                "frame": 0, "x": 50, "y": 60, "ancho": 300, "alto": 150,
                "giro": 0, "opacidad": 1, "color": "#ffffff", "tamano": 36,
                "texto": "Inicio", "fuente": "Georgia", "interpolacion": "bezier",
                "bezier": [0.2, -0.2, 0.8, 1.2], "texto_modo": "instantaneo",
            },
            {
                "frame": 9000, "x": 700, "y": -100, "ancho": 420, "alto": 180,
                "giro": 180, "opacidad": 0.4, "color": "#123456", "tamano": 60,
                "texto": "Final", "fuente": "Courier New", "interpolacion": "suave",
                "texto_modo": "escritura",
            },
        ]
        audio_path = "audio_uploads/1/0123456789abcdef0123456789abcdef_pista.wav"
        tracks = [
            {
                "id": f"track-{index}", "path": audio_path, "inicio_frame": index * 30,
                "recorte_inicio": 0, "recorte_fin": 12.5, "volumen": 0.8,
                "silenciado": index == 2, "fundido_entrada": 0.5, "fundido_salida": 1,
            }
            for index in range(3)
        ]
        slide = {
            "id": "slide-1", "objetos": [objeto], "fondo": "", "color_fondo": "#ffffff",
            "audio": audio_path, "audios": tracks,
            "marcadores": [{"id": "marker-1", "frame": 4500, "objeto_id": objeto["id"]}],
            "duracion_frames": 9000,
        }
        validar_diapositivas([slide])
        with self.assertRaises(ValidationError):
            validar_diapositivas([{**slide, "audios": tracks + [{**tracks[0], "id": "track-4"}]}])
        with self.assertRaises(ValidationError):
            validar_diapositivas([{**slide, "marcadores": [{"id": "bad", "frame": 2, "objeto_id": "missing"}]}])

    def test_evaluador_de_exportacion_interpola_propiedades_y_texto(self):
        objeto = self.cuadro()
        objeto["fotogramas"] = [
            {"frame": 0, "x": 0, "y": 0, "giro": 0, "opacidad": 1, "interpolacion": "lineal"},
            {
                "frame": 30, "x": 300, "y": 120, "ancho": 600, "alto": 300,
                "giro": 90, "opacidad": 0.5, "color": "#000000", "tamano": 72,
                "texto": "Animado", "fuente": "Arial", "interpolacion": "lineal",
                "texto_modo": "escritura",
            },
        ]
        pose = _evaluate_object(objeto, 15)
        self.assertEqual((pose["x"], pose["y"], pose["giro"]), (150, 60, 45))
        self.assertEqual((pose["ancho"], pose["alto"], pose["tamano"]), (450, 225, 54))
        self.assertEqual(pose["texto"], "Anim")

    def test_plantillas_de_animacion_son_privadas_editables_y_eliminables(self):
        list_url = reverse("editor:plantillas_animacion")
        save_url = reverse("editor:guardar_plantilla_animacion")
        data = {"fotogramas": [
            {
                "frame": 0, "x": 0, "y": 0, "ancho": 100, "alto": 100,
                "giro": 0, "opacidad": 0, "color": "#ffffff", "tamano": 36,
                "texto": "", "fuente": "Poppins", "interpolacion": "suave",
                "texto_modo": "instantaneo",
            },
            {
                "frame": 20, "x": 20, "y": 10, "ancho": 100, "alto": 100,
                "giro": 0, "opacidad": 1, "color": "#ffffff", "tamano": 36,
                "texto": "", "fuente": "Poppins", "interpolacion": "suave",
                "texto_modo": "instantaneo",
            },
        ]}
        response = self.client.post(save_url, {"nombre": "Entrada propia", "datos": json.dumps(data)})
        self.assertEqual(response.status_code, 201)
        template_id = response.json()["template"]["id"]
        self.assertEqual(PlantillaAnimacion.objects.filter(propietario=self.usuario).count(), 1)
        self.assertEqual(self.client.get(list_url).json()["templates"][0]["data"], data)

        self.client.force_login(self.otro)
        self.assertEqual(self.client.get(list_url).json()["templates"], [])
        self.assertEqual(
            self.client.post(reverse("editor:eliminar_plantilla_animacion", args=[template_id])).status_code,
            404,
        )
        self.client.force_login(self.usuario)
        self.assertEqual(
            self.client.post(reverse("editor:eliminar_plantilla_animacion", args=[template_id])).status_code,
            200,
        )
        self.assertFalse(PlantillaAnimacion.objects.filter(pk=template_id).exists())

    def test_exportacion_mp4_y_gif_acepta_solo_1080_y_720(self):
        with patch("editor.views.exportar_video", return_value=b"mp4") as video:
            response = self.client.get(
                reverse("editor:exportar", args=[self.proyecto.pk, "mp4"]),
                {"resolucion": "720"},
            )
            self.assertEqual((response.status_code, response["Content-Type"], response.content), (200, "video/mp4", b"mp4"))
            video.assert_called_once()
            self.assertEqual(video.call_args.args[1], 720)
        with patch("editor.views.exportar_gif", return_value=b"GIF89a") as gif:
            response = self.client.get(reverse("editor:exportar", args=[self.proyecto.pk, "gif"]))
            self.assertEqual((response.status_code, response["Content-Type"]), (200, "image/gif"))
            gif.assert_called_once()
            self.assertEqual(gif.call_args.args[1], 1080)
        self.assertEqual(
            self.client.get(reverse("editor:exportar", args=[self.proyecto.pk, "mp4"]), {"resolucion": "4k"}).status_code,
            400,
        )

    def test_exportacion_mp4_mezcla_audio_de_las_pistas(self):
        audio = BytesIO()
        with wave.open(audio, "wb") as wav:
            wav.setnchannels(1)
            wav.setsampwidth(2)
            wav.setframerate(8000)
            wav.writeframes(b"\x00\x00" * 800)
        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            path = f"audio_uploads/{self.usuario.pk}/0123456789abcdef0123456789abcdef_tono.wav"
            default_storage.save(path, ContentFile(audio.getvalue()))
            slide = {
                "id": "slide-1", "objetos": [], "fondo": "", "color_fondo": "#ffffff",
                "duracion_frames": 3, "audio": path,
                "audios": [{
                    "id": "track-1", "path": path, "inicio_frame": 0,
                    "recorte_inicio": 0, "recorte_fin": 0.1, "volumen": 0.7,
                    "silenciado": False, "fundido_entrada": 0.02, "fundido_salida": 0.02,
                }],
                "marcadores": [],
            }
            video = exportar_video([slide], 720)
        self.assertEqual(video[4:8], b"ftyp")
        self.assertIn(b"mp4a", video)

    @override_settings(GEMINI_API_KEY="clave-super-secreta")
    def test_panel_ia_muestra_estado_sin_exponer_clave(self):
        editor = self.client.get(self.detalle)
        self.assertContains(editor, 'id="ai-tool"')
        self.assertContains(editor, self.estado_ia)
        self.assertContains(editor, self.generar_ia)
        self.assertNotContains(editor, "clave-super-secreta")

        state = self.client.get(self.estado_ia)
        self.assertEqual(state.status_code, 200)
        self.assertTrue(state.json()["configured"])
        self.assertNotIn("api_key", state.json())
        self.assertEqual(state["Cache-Control"], "private, no-store")

    @override_settings(GEMINI_API_KEY="")
    def test_ia_sin_clave_no_llama_servicio(self):
        with patch("editor.views.gemini_service.generate_presentation") as generate:
            response = self.client.post(
                self.generar_ia,
                data=json.dumps({
                    "action": "presentation", "topic": "Biología", "slide_count": 3,
                    "audience": "Bachillerato", "style": "academico",
                }),
                content_type="application/json",
            )
        self.assertEqual(response.status_code, 503)
        self.assertIn("GEMINI_API_KEY", response.json()["error"])
        generate.assert_not_called()
        self.assertFalse(UsoIA.objects.exists())

    @override_settings(GEMINI_API_KEY="prueba", GEMINI_DAILY_TOKEN_LIMIT=1000)
    def test_ia_genera_presentacion_registra_consumo_y_respeta_propiedad(self):
        generated = {
            "title": "Energía solar",
            "slides": [
                {
                    "title": "Introducción",
                    "bullets": ["Concepto", "Aplicaciones"],
                    "speaker_notes": "Explica el contexto.",
                    "visual_suggestion": "Diagrama del Sol y un panel.",
                },
                {
                    "title": "Conclusiones",
                    "bullets": ["Beneficios", "Retos"],
                    "speaker_notes": "Cierra con una pregunta.",
                    "visual_suggestion": "Tabla comparativa.",
                },
            ],
        }
        result = GeminiResult(generated, TokenUsage(120, 80, 200))
        with patch("editor.views.gemini_service.generate_presentation", return_value=result) as generate:
            response = self.client.post(
                self.generar_ia,
                data=json.dumps({
                    "action": "presentation", "topic": "Energía solar", "slide_count": 2,
                    "audience": "Universidad", "style": "visual",
                }),
                content_type="application/json",
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["presentation"], generated)
        self.assertEqual(response.json()["state"]["quota"]["used"], 200)
        generate.assert_called_once_with("Energía solar", 2, "Universidad", "visual")
        usage = UsoIA.objects.get()
        self.assertEqual((usage.accion, usage.tokens_totales, usage.proyecto), ("presentacion", 200, self.proyecto))

        foreign = reverse("editor:generar_ia", args=[self.ajeno.pk])
        self.assertEqual(self.client.post(foreign, data="{}", content_type="application/json").status_code, 404)

    @override_settings(GEMINI_API_KEY="prueba", GEMINI_DAILY_TOKEN_LIMIT=10)
    def test_ia_bloquea_al_alcanzar_cuota(self):
        UsoIA.objects.create(
            usuario=self.usuario, proyecto=self.proyecto, accion="texto", modelo="modelo",
            tokens_totales=10,
        )
        with patch("editor.views.gemini_service.transform_text") as transform:
            response = self.client.post(
                self.generar_ia,
                data=json.dumps({"action": "text", "text_action": "resumir", "source": "Texto largo"}),
                content_type="application/json",
            )
        self.assertEqual(response.status_code, 429)
        transform.assert_not_called()

    @override_settings(GEMINI_API_KEY="prueba")
    def test_ia_texto_e_imagen_se_validan_registran_y_guardan_en_biblioteca_privada(self):
        text_result = GeminiResult("Resumen claro", TokenUsage(20, 5, 25))
        with patch("editor.views.gemini_service.transform_text", return_value=text_result) as transform:
            response = self.client.post(
                self.generar_ia,
                data=json.dumps({
                    "action": "text", "text_action": "resumir", "source": "Contenido original",
                    "instruction": "Breve",
                }),
                content_type="application/json",
            )
        self.assertEqual((response.status_code, response.json()["text"]), (200, "Resumen claro"))
        transform.assert_called_once_with("resumir", "Contenido original", "Breve")

        image_buffer = BytesIO()
        Image.new("RGB", (64, 36), "blue").save(image_buffer, format="PNG")
        image_result = GeminiResult(
            {"bytes": image_buffer.getvalue(), "mime_type": "image/png"},
            TokenUsage(30, 40, 70),
        )
        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            with patch("editor.views.gemini_service.generate_image", return_value=image_result) as generate:
                response = self.client.post(
                    self.generar_ia,
                    data=json.dumps({
                        "action": "image", "prompt": "Ciclo del agua", "image_type": "diagrama",
                        "aspect_ratio": "16:9",
                    }),
                    content_type="application/json",
                )
            self.assertEqual(response.status_code, 201)
            image_data = response.json()["image"]
            self.assertTrue(image_data["path"].startswith(f"uploads/{self.usuario.pk}/"))
            self.assertTrue(default_storage.exists(image_data["path"]))
            image_response = self.client.get(image_data["url"])
            self.assertEqual(image_response.status_code, 200)
            image_response.close()
            self.client.force_login(self.otro)
            self.assertEqual(self.client.get(image_data["url"]).status_code, 404)
        generate.assert_called_once_with("Ciclo del agua", "diagrama", "16:9")
        self.assertEqual(UsoIA.objects.filter(usuario=self.usuario, exitoso=True).count(), 2)

    @override_settings(GEMINI_API_KEY="prueba")
    def test_ia_error_de_servicio_no_filtra_detalles_y_registra_fallo(self):
        with patch(
            "editor.views.gemini_service.transform_text",
            side_effect=GeminiServiceError("Gemini no pudo procesar el texto. Intenta de nuevo."),
        ):
            response = self.client.post(
                self.generar_ia,
                data=json.dumps({"action": "text", "text_action": "mejorar", "source": "Texto"}),
                content_type="application/json",
            )
        self.assertEqual(response.status_code, 502)
        usage = UsoIA.objects.get()
        self.assertFalse(usage.exitoso)
        self.assertEqual(usage.tokens_totales, 0)

    def _pptx_upload(self, filename="clase.pptx"):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Energía renovable"
        slide.placeholders[1].text = "Solar\nEólica"
        buffer = BytesIO()
        presentation.save(buffer)
        return SimpleUploadedFile(
            filename,
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )

    @override_settings(GEMINI_API_KEY="prueba")
    def test_conversion_documento_valida_crea_trabajo_privado_y_permite_cancelar_reintentar(self):
        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            with patch("editor.views.convertir_documento.delay", return_value=SimpleNamespace(id="task-1")):
                response = self.client.post(self.subir_documento_ia, {
                    "document": self._pptx_upload(),
                    "slide_count": "4",
                    "audience": "Licenciatura",
                    "style": "visual",
                })
            self.assertEqual(response.status_code, 202)
            job = TrabajoConversion.objects.select_related("documento").get()
            self.assertEqual(job.documento.tipo, DocumentoFuente.Tipo.PPTX)
            self.assertEqual(job.numero_diapositivas, 4)
            self.assertTrue(default_storage.exists(job.documento.archivo.name))
            self.assertNotIn(job.documento.archivo.name, response.content.decode())

            status_url = reverse("editor:estado_conversion", args=[self.proyecto.pk, job.pk])
            cancel_url = reverse("editor:cancelar_conversion", args=[self.proyecto.pk, job.pk])
            retry_url = reverse("editor:reintentar_conversion", args=[self.proyecto.pk, job.pk])
            self.assertEqual(self.client.get(status_url).json()["conversion"]["status"], "pendiente")
            self.assertEqual(self.client.post(cancel_url).json()["conversion"]["status"], "cancelado")
            with patch("editor.views.convertir_documento.delay", return_value=SimpleNamespace(id="task-2")):
                retried = self.client.post(retry_url)
            self.assertEqual(retried.status_code, 202)
            self.assertEqual(retried.json()["conversion"]["attempts"], 2)

            self.client.force_login(self.otro)
            self.assertEqual(self.client.get(status_url).status_code, 404)

    @override_settings(GEMINI_API_KEY="prueba")
    def test_conversion_rechaza_archivos_disfrazados_y_opciones_invalidas(self):
        fake_pdf = SimpleUploadedFile("falso.pdf", b"no-es-pdf", content_type="application/pdf")
        response = self.client.post(self.subir_documento_ia, {
            "document": fake_pdf,
            "slide_count": "5",
            "audience": "Universidad",
            "style": "academico",
        })
        self.assertEqual(response.status_code, 400)
        self.assertFalse(DocumentoFuente.objects.exists())

        response = self.client.post(self.subir_documento_ia, {
            "document": self._pptx_upload(),
            "slide_count": "21",
            "audience": "Universidad",
            "style": "academico",
        })
        self.assertEqual(response.status_code, 400)
        self.assertFalse(DocumentoFuente.objects.exists())

    @override_settings(GEMINI_API_KEY="prueba", GEMINI_DAILY_TOKEN_LIMIT=1000)
    def test_tarea_pdf_usa_analisis_visual_y_el_resultado_no_expone_archivo(self):
        pdf = BytesIO()
        writer = canvas.Canvas(pdf)
        writer.drawString(72, 720, "Contenido visual del documento")
        writer.save()
        upload = SimpleUploadedFile("lectura.pdf", pdf.getvalue(), content_type="application/pdf")
        generated = {
            "title": "Lectura",
            "slides": [{
                "title": "Idea principal",
                "bullets": ["Contenido del documento"],
                "speaker_notes": "Explica la idea.",
                "visual_suggestion": "Conservar el esquema principal.",
            }],
        }
        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            with patch("editor.views.convertir_documento.delay", return_value=SimpleNamespace(id="pdf-task")):
                response = self.client.post(self.subir_documento_ia, {
                    "document": upload,
                    "slide_count": "1",
                    "audience": "Universidad",
                    "style": "academico",
                })
            self.assertEqual(response.status_code, 202)
            job = TrabajoConversion.objects.select_related("documento").get()
            self.assertEqual(job.documento.tipo, DocumentoFuente.Tipo.PDF)
            service_result = GeminiResult(generated, TokenUsage(260, 30, 290))
            with patch(
                "editor.tasks.gemini_service.generate_presentation_from_pdf",
                return_value=service_result,
            ) as generate:
                result = convertir_documento.run(str(job.pk))
            job.refresh_from_db()
            self.assertEqual(result["state"], "completed")
            self.assertEqual(job.resultado, generated)
            self.assertTrue(hasattr(generate.call_args.args[0], "read"))
            status = self.client.get(reverse("editor:estado_conversion", args=[self.proyecto.pk, job.pk]))
            self.assertNotIn(job.documento.archivo.name, status.content.decode())

    @override_settings(GEMINI_API_KEY="prueba", GEMINI_DAILY_TOKEN_LIMIT=1000)
    def test_tarea_convierte_pptx_registra_consumo_y_guarda_resultado(self):
        generated = {
            "title": "Energía",
            "slides": [{
                "title": "Fuentes renovables",
                "bullets": ["Solar", "Eólica"],
                "speaker_notes": "Explica las diferencias.",
                "visual_suggestion": "Dos iconos comparativos.",
            }],
        }
        with tempfile.TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            document = DocumentoFuente.objects.create(
                usuario=self.usuario,
                proyecto=self.proyecto,
                archivo=self._pptx_upload(),
                nombre_original="clase.pptx",
                tipo=DocumentoFuente.Tipo.PPTX,
                tamano=100,
            )
            job = TrabajoConversion.objects.create(
                usuario=self.usuario,
                proyecto=self.proyecto,
                documento=document,
                numero_diapositivas=1,
                audiencia="Universidad",
                estilo="academico",
            )
            service_result = GeminiResult(generated, TokenUsage(80, 40, 120))
            with patch(
                "editor.tasks.gemini_service.generate_presentation_from_extracted_document",
                return_value=service_result,
            ) as generate:
                task_result = convertir_documento.run(str(job.pk))
            job.refresh_from_db()
        self.assertEqual(task_result["state"], "completed")
        self.assertEqual((job.estado, job.progreso, job.resultado), ("completado", 100, generated))
        self.assertIn("Energía renovable", generate.call_args.args[0])
        usage = UsoIA.objects.get(accion=UsoIA.Accion.DOCUMENTO)
        self.assertEqual((usage.tokens_totales, usage.exitoso), (120, True))
