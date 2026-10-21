from contextlib import contextmanager
from io import BytesIO
import shutil
import subprocess
import tempfile
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from django.core.files.storage import default_storage
from PIL import Image, ImageColor, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor

from .scene import IMAGE_DIRECTORY


WIDTH, HEIGHT = 800, 450
EXPORT_WIDTH, EXPORT_HEIGHT = 1600, 900
ANIMATION_FPS = 30
PPT_WIDTH, PPT_HEIGHT = 13.333333, 7.5


@contextmanager
def _open_image(relative_path):
    if relative_path.startswith("uploads/"):
        with default_storage.open(relative_path, "rb") as image_file:
            yield image_file
    else:
        with (IMAGE_DIRECTORY / relative_path).open("rb") as image_file:
            yield image_file


def _image_bytes(relative_path):
    with _open_image(relative_path) as image_file, Image.open(image_file) as image:
        output = BytesIO()
        image.convert("RGBA").save(output, format="PNG")
    return output.getvalue()


def _font_path(bold=False, italic=False, family="Arial"):
    windows_names = {
        "Poppins": {
            (False, False): "Poppins-Regular.ttf", (True, False): "Poppins-Bold.ttf",
            (False, True): "Poppins-Italic.ttf", (True, True): "Poppins-BoldItalic.ttf",
        },
        "Arial": {
            (False, False): "arial.ttf", (True, False): "arialbd.ttf",
            (False, True): "ariali.ttf", (True, True): "arialbi.ttf",
        },
        "Georgia": {
            (False, False): "georgia.ttf", (True, False): "georgiab.ttf",
            (False, True): "georgiai.ttf", (True, True): "georgiaz.ttf",
        },
        "Verdana": {
            (False, False): "verdana.ttf", (True, False): "verdanab.ttf",
            (False, True): "verdanai.ttf", (True, True): "verdanaz.ttf",
        },
        "Courier New": {
            (False, False): "cour.ttf", (True, False): "courbd.ttf",
            (False, True): "couri.ttf", (True, True): "courbi.ttf",
        },
    }
    name = windows_names.get(family, windows_names["Arial"])[(bold, italic)]
    linux_name = {
        (False, False): ("DejaVuSans.ttf", "LiberationSans-Regular.ttf"),
        (True, False): ("DejaVuSans-Bold.ttf", "LiberationSans-Bold.ttf"),
        (False, True): ("DejaVuSans-Oblique.ttf", "LiberationSans-Italic.ttf"),
        (True, True): ("DejaVuSans-BoldOblique.ttf", "LiberationSans-BoldItalic.ttf"),
    }[(bold, italic)]
    candidates = [
        Path("C:/Windows/Fonts") / name,
        Path("C:/Windows/Fonts") / windows_names["Arial"][(bold, italic)],
        Path("/usr/share/fonts/truetype/dejavu") / linux_name[0],
        Path("/usr/share/fonts/truetype/liberation2") / linux_name[1],
    ]
    return next((path for path in candidates if path.is_file()), None)


def _load_pil_font(size, bold=False, italic=False, family="Arial"):
    path = _font_path(bold, italic, family)
    return ImageFont.truetype(str(path), size) if path else ImageFont.load_default(size=size)


def _wrap_lines(text, max_width, measure):
    output = []
    for paragraph in text.splitlines() or [""]:
        line = ""
        for char in paragraph:
            if line and measure(line + char) > max_width:
                output.append(line)
                line = ""
            line += char
        output.append(line)
    return output


def _pdf_font(bold, italic):
    name = f"EditorFont{int(bold)}{int(italic)}"
    if name not in pdfmetrics.getRegisteredFontNames():
        path = _font_path(bold, italic)
        if path:
            pdfmetrics.registerFont(TTFont(name, str(path)))
        else:
            return "Helvetica-BoldOblique" if bold and italic else (
                "Helvetica-Bold" if bold else "Helvetica-Oblique" if italic else "Helvetica"
            )
    return name


def _cover_background(canvas_image, relative_path):
    if not relative_path:
        return
    with _open_image(relative_path) as image_file, Image.open(image_file) as background:
        fitted = ImageOps.fit(background.convert("RGB"), canvas_image.size)
    canvas_image.paste(fitted)


def _slide_color(slide_data):
    return slide_data.get("color_fondo", "#ffffff")


def _object_rotation(item):
    return next(
        (frame["giro"] for frame in item["fotogramas"] if frame["frame"] == 0),
        0,
    )


def _cubic_bezier_value(x, values):
    x1, y1, x2, y2 = values

    def sample(value, first, second):
        return 3 * (1 - value) ** 2 * value * first + 3 * (1 - value) * value ** 2 * second + value ** 3

    low, high, value = 0.0, 1.0, x
    for _ in range(12):
        current_x = sample(value, x1, x2)
        if abs(current_x - x) < 0.0005:
            break
        if current_x < x:
            low = value
        else:
            high = value
        value = (low + high) / 2
    return sample(value, y1, y2)


def _eased_progress(progress, keyframe):
    interpolation = keyframe.get("interpolacion", "lineal")
    if interpolation == "instantanea":
        return 0 if progress < 1 else 1
    if interpolation == "suave":
        return progress * progress * (3 - 2 * progress)
    if interpolation == "rebote":
        n1, d1 = 7.5625, 2.75
        if progress < 1 / d1:
            return n1 * progress * progress
        if progress < 2 / d1:
            progress -= 1.5 / d1
            return n1 * progress * progress + 0.75
        if progress < 2.5 / d1:
            progress -= 2.25 / d1
            return n1 * progress * progress + 0.9375
        progress -= 2.625 / d1
        return n1 * progress * progress + 0.984375
    if interpolation == "bezier":
        return _cubic_bezier_value(progress, keyframe.get("bezier", [0.25, 0.1, 0.25, 1]))
    return progress


def _normalized_keyframes(item):
    pose = {
        "x": item["x"], "y": item["y"], "ancho": item["ancho"], "alto": item["alto"],
        "giro": 0, "opacidad": 1, "color": item["color"], "tamano": item["tamano"],
        "texto": item["texto"], "fuente": item["fuente"], "texto_modo": "instantaneo",
    }
    output = []
    for source in sorted(item["fotogramas"], key=lambda key: key["frame"]):
        pose = {**pose, **source}
        output.append(dict(pose))
    return output


def _interpolate_color(left, right, progress):
    left_rgb = ImageColor.getrgb(left)
    right_rgb = ImageColor.getrgb(right)
    rgb = tuple(round(a + (b - a) * progress) for a, b in zip(left_rgb, right_rgb))
    return "#" + "".join(f"{value:02x}" for value in rgb)


def _evaluate_object(item, frame):
    keyframes = _normalized_keyframes(item)
    if frame <= keyframes[0]["frame"]:
        return {**keyframes[0], "text_opacity": 1}
    if frame >= keyframes[-1]["frame"]:
        return {**keyframes[-1], "text_opacity": 1}
    right_index = next(index for index, key in enumerate(keyframes) if key["frame"] > frame)
    left, right = keyframes[right_index - 1], keyframes[right_index]
    raw_progress = (frame - left["frame"]) / (right["frame"] - left["frame"])
    progress = _eased_progress(raw_progress, left)
    pose = dict(left)
    for property_name in ("x", "y", "ancho", "alto", "giro", "opacidad", "tamano"):
        pose[property_name] = left[property_name] + (right[property_name] - left[property_name]) * progress
    pose["color"] = _interpolate_color(left["color"], right["color"], max(0, min(1, progress)))
    pose["fuente"] = left["fuente"] if raw_progress < 0.5 else right["fuente"]
    text_mode = right.get("texto_modo", "instantaneo")
    if text_mode == "escritura":
        pose["texto"] = right["texto"][:round(len(right["texto"]) * max(0, min(1, progress)))]
        pose["text_opacity"] = 1
    elif text_mode == "transicion":
        pose["texto"] = left["texto"] if raw_progress < 0.5 else right["texto"]
        pose["text_opacity"] = abs(raw_progress - 0.5) * 2
    else:
        pose["texto"] = left["texto"]
        pose["text_opacity"] = 1
    return pose


def _composite_rotated(canvas_image, layer, x, y, width, height, rotation):
    if rotation:
        layer = layer.rotate(-rotation, expand=True, resample=Image.Resampling.BICUBIC)
    left = round(x + (width - layer.width) / 2)
    top = round(y + (height - layer.height) / 2)
    source_left = max(0, -left)
    source_top = max(0, -top)
    source_right = min(layer.width, canvas_image.width - left)
    source_bottom = min(layer.height, canvas_image.height - top)
    if source_left >= source_right or source_top >= source_bottom:
        return
    canvas_image.alpha_composite(
        layer.crop((source_left, source_top, source_right, source_bottom)),
        (max(0, left), max(0, top)),
    )


def _cover_image_bytes(relative_path):
    with _open_image(relative_path) as image_file, Image.open(image_file) as image:
        fitted = ImageOps.fit(image.convert("RGB"), (EXPORT_WIDTH, EXPORT_HEIGHT))
        output = BytesIO()
        fitted.save(output, format="PNG")
    return output.getvalue()


def _render_slide(slide_data, frame=None, output_size=(EXPORT_WIDTH, EXPORT_HEIGHT)):
    output_width, output_height = output_size
    image = Image.new("RGBA", output_size, ImageColor.getrgb(_slide_color(slide_data)))
    _cover_background(image, slide_data["fondo"])
    scale = output_width / WIDTH
    for item in slide_data["objetos"]:
        pose = _evaluate_object(item, frame) if frame is not None else {
            "x": item["x"], "y": item["y"], "ancho": item["ancho"], "alto": item["alto"],
            "giro": _object_rotation(item), "opacidad": 1, "color": item["color"],
            "tamano": item["tamano"], "texto": item["texto"], "fuente": item["fuente"],
            "text_opacity": 1,
        }
        x, y = round(pose["x"] * scale), round(pose["y"] * scale)
        width, height = max(1, round(pose["ancho"] * scale)), max(1, round(pose["alto"] * scale))
        layer = Image.new("RGBA", (width, height), (0, 0, 0, 0))
        if item.get("tipo") == "imagen":
            with _open_image(item["imagen"]) as image_file, Image.open(image_file) as source:
                source = source.convert("RGBA")
                source.thumbnail((width, height), Image.Resampling.LANCZOS)
                layer.alpha_composite(source, ((width - source.width) // 2, (height - source.height) // 2))
        else:
            draw = ImageDraw.Draw(layer)
            font_size = max(1, round(pose["tamano"] * scale))
            font = _load_pil_font(font_size, item["negrita"], item["cursiva"], pose["fuente"])
            line_height = font_size * 1.2
            lines = _wrap_lines(pose["texto"], width - 8 * scale, lambda line: draw.textlength(line, font=font))
            text_alpha = round(255 * pose.get("text_opacity", 1))
            text_color = (*ImageColor.getrgb(pose["color"]), text_alpha)
            for index, line in enumerate(lines[:max(1, int(height / line_height))]):
                line_width = draw.textlength(line, font=font)
                text_x = {"left": 0, "center": (width - line_width) / 2, "right": width - line_width}[item["alineacion"]]
                text_y = 4 * scale + index * line_height
                draw.text((text_x, text_y), line, font=font, fill=text_color)
                if item["subrayado"]:
                    draw.line((text_x, text_y + font_size, text_x + line_width, text_y + font_size), fill=text_color, width=max(1, round(scale)))
        if pose["opacidad"] < 1:
            alpha = layer.getchannel("A").point(lambda value: round(value * pose["opacidad"]))
            layer.putalpha(alpha)
        _composite_rotated(image, layer, x, y, width, height, pose["giro"])
    return image.convert("RGB")


def exportar_imagenes(diapositivas):
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for index, diapositiva in enumerate(diapositivas, start=1):
            buffer = BytesIO()
            _render_slide(diapositiva).save(buffer, format="PNG")
            archive.writestr(f"diapositiva-{index}.png", buffer.getvalue())
    return output.getvalue()


def exportar_pdf(nombre, diapositivas):
    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=(PPT_WIDTH * 72, PPT_HEIGHT * 72), pageCompression=1)
    scale = PPT_WIDTH * 72 / WIDTH
    for diapositiva in diapositivas:
        pdf.setFillColor(HexColor(_slide_color(diapositiva)))
        pdf.rect(0, 0, PPT_WIDTH * 72, PPT_HEIGHT * 72, fill=1, stroke=0)
        if diapositiva["fondo"]:
            pdf.drawImage(
                ImageReader(BytesIO(_cover_image_bytes(diapositiva["fondo"]))),
                0, 0, width=PPT_WIDTH * 72, height=PPT_HEIGHT * 72,
                preserveAspectRatio=False, mask="auto",
            )
        for item in diapositiva["objetos"]:
            x, y = item["x"] * scale, (HEIGHT - item["y"] - item["alto"]) * scale
            rotation = _object_rotation(item)
            pdf.saveState()
            pdf.translate(x + item["ancho"] * scale / 2, y + item["alto"] * scale / 2)
            pdf.rotate(-rotation)
            if item.get("tipo") == "imagen":
                pdf.drawImage(
                    ImageReader(BytesIO(_image_bytes(item["imagen"]))),
                    -item["ancho"] * scale / 2, -item["alto"] * scale / 2,
                    width=item["ancho"] * scale, height=item["alto"] * scale,
                    preserveAspectRatio=True, anchor="c", mask="auto",
                )
                pdf.restoreState()
                continue
            font_name = _pdf_font(item["negrita"], item["cursiva"])
            pdf.setFont(font_name, item["tamano"] * scale)
            pdf.setFillColor(item["color"])
            lines = _wrap_lines(
                item["texto"],
                item["ancho"] * scale * 72 - 6,
                lambda line: pdfmetrics.stringWidth(line, font_name, item["tamano"] * scale),
            )
            for line_index, line in enumerate(lines[:max(1, int(item["alto"] / (item["tamano"] * 1.2)))]):
                line_y = item["alto"] * scale / 2 - (line_index + 1) * item["tamano"] * 1.2 * scale
                if item["alineacion"] == "center":
                    pdf.drawCentredString(0, line_y, line)
                elif item["alineacion"] == "right":
                    pdf.drawRightString(item["ancho"] * scale / 2, line_y, line)
                else:
                    pdf.drawString(-item["ancho"] * scale / 2, line_y, line)
            pdf.restoreState()
        pdf.showPage()
    pdf.setTitle(nombre)
    pdf.save()
    return output.getvalue()


def _add_picture(slide, relative_path, x, y, width, height):
    with _open_image(relative_path) as image_file, Image.open(image_file) as source:
        image_width, image_height = source.size
    scale = min(width / image_width, height / image_height)
    fitted_width, fitted_height = width * scale, height * scale
    return slide.shapes.add_picture(
        BytesIO(_image_bytes(relative_path)),
        Inches(x + (width - fitted_width) / 2),
        Inches(y + (height - fitted_height) / 2),
        width=Inches(fitted_width),
        height=Inches(fitted_height),
    )


def exportar_pptx(nombre, diapositivas):
    presentation = Presentation()
    presentation.slide_width = Inches(PPT_WIDTH)
    presentation.slide_height = Inches(PPT_HEIGHT)
    blank_layout = presentation.slide_layouts[6]
    scale = PPT_WIDTH / WIDTH
    alignments = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
    for data in diapositivas:
        slide = presentation.slides.add_slide(blank_layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(_slide_color(data).lstrip("#"))
        if data["fondo"]:
            slide.shapes.add_picture(
                BytesIO(_cover_image_bytes(data["fondo"])), 0, 0,
                width=presentation.slide_width, height=presentation.slide_height,
            )
        for item in data["objetos"]:
            x, y = item["x"] * scale, item["y"] * scale
            width, height = item["ancho"] * scale, item["alto"] * scale
            if item.get("tipo") == "imagen":
                picture = _add_picture(slide, item["imagen"], x, y, width, height)
                picture.rotation = _object_rotation(item)
                continue
            shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
            shape.rotation = _object_rotation(item)
            frame = shape.text_frame
            frame.clear()
            frame.word_wrap = True
            frame.margin_left = Pt(3)
            frame.margin_right = Pt(3)
            frame.margin_top = Pt(3)
            frame.margin_bottom = Pt(3)
            frame.vertical_anchor = MSO_ANCHOR.TOP
            for line_index, line in enumerate(item["texto"].splitlines() or [""]):
                paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
                paragraph.text = line
                paragraph.alignment = alignments[item["alineacion"]]
                paragraph.font.name = item["fuente"]
                paragraph.font.size = Pt(item["tamano"] * scale * 72)
                paragraph.font.bold = item["negrita"]
                paragraph.font.italic = item["cursiva"]
                paragraph.font.underline = item["subrayado"]
                paragraph.font.color.rgb = RGBColor.from_string(item["color"].lstrip("#"))
    presentation.core_properties.title = nombre
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def _ffmpeg_executable():
    try:
        import imageio_ffmpeg
    except ImportError as error:
        raise RuntimeError("La exportación animada no está instalada.") from error
    return imageio_ffmpeg.get_ffmpeg_exe()


def _run_ffmpeg(command):
    result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=False)
    if result.returncode:
        message = result.stderr.decode("utf-8", errors="replace")[-2000:]
        raise RuntimeError(f"FFmpeg no pudo completar la exportación: {message}")


def _write_animation_frames(diapositivas, output_path, output_size, gif=False):
    width, height = output_size
    command = [
        _ffmpeg_executable(), "-y", "-loglevel", "error", "-f", "rawvideo",
        "-pix_fmt", "rgb24", "-s", f"{width}x{height}", "-r", str(ANIMATION_FPS),
        "-i", "pipe:0", "-an",
    ]
    if gif:
        command += ["-vf", "split[s0][s1];[s0]palettegen=max_colors=256[p];[s1][p]paletteuse=dither=sierra2_4a", "-loop", "0"]
    else:
        command += ["-c:v", "libx264", "-preset", "medium", "-crf", "20", "-pix_fmt", "yuv420p", "-movflags", "+faststart"]
    command.append(str(output_path))
    process = subprocess.Popen(command, stdin=subprocess.PIPE, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
    try:
        for slide in diapositivas:
            for frame in range(slide["duracion_frames"]):
                process.stdin.write(_render_slide(slide, frame, output_size).tobytes())
        process.stdin.close()
        error_output = process.stderr.read()
        return_code = process.wait()
    except (BrokenPipeError, OSError):
        process.kill()
        error_output = process.stderr.read()
        process.wait()
        raise RuntimeError(error_output.decode("utf-8", errors="replace")[-2000:])
    if return_code:
        raise RuntimeError(error_output.decode("utf-8", errors="replace")[-2000:])


def _audio_sources(diapositivas):
    slide_offset = 0
    sources = []
    for slide in diapositivas:
        tracks = slide.get("audios") or []
        if not tracks and slide.get("audio"):
            tracks = [{
                "path": slide["audio"], "inicio_frame": 0, "recorte_inicio": 0,
                "recorte_fin": None, "volumen": 1, "silenciado": False,
                "fundido_entrada": 0, "fundido_salida": 0,
            }]
        for track in tracks:
            if not track.get("silenciado"):
                sources.append((slide_offset, track))
        slide_offset += slide["duracion_frames"] / ANIMATION_FPS
    return sources, slide_offset


def _mix_audio(video_path, output_path, diapositivas, temporary_directory):
    sources, total_duration = _audio_sources(diapositivas)
    if not sources:
        shutil.copyfile(video_path, output_path)
        return
    command = [_ffmpeg_executable(), "-y", "-loglevel", "error", "-i", str(video_path)]
    filters = []
    labels = []
    for index, (slide_offset, track) in enumerate(sources, start=1):
        suffix = Path(track["path"]).suffix or ".audio"
        audio_path = Path(temporary_directory) / f"audio-{index}{suffix}"
        with default_storage.open(track["path"], "rb") as source, audio_path.open("wb") as destination:
            shutil.copyfileobj(source, destination)
        command += ["-i", str(audio_path)]
        trim_start = track.get("recorte_inicio", 0)
        trim_end = track.get("recorte_fin")
        trim = f"atrim=start={trim_start}"
        if trim_end is not None:
            trim += f":end={trim_end}"
        duration = None if trim_end is None else max(0, trim_end - trim_start)
        chain = [trim, "asetpts=PTS-STARTPTS", f"volume={track.get('volumen', 1)}"]
        fade_in = track.get("fundido_entrada", 0)
        fade_out = track.get("fundido_salida", 0)
        if fade_in:
            chain.append(f"afade=t=in:st=0:d={fade_in}")
        if fade_out and duration is not None:
            chain.append(f"afade=t=out:st={max(0, duration - fade_out)}:d={min(fade_out, duration)}")
        delay = round((slide_offset + track.get("inicio_frame", 0) / ANIMATION_FPS) * 1000)
        chain.append(f"adelay={delay}:all=1")
        label = f"a{index}"
        filters.append(f"[{index}:a]{','.join(chain)}[{label}]")
        labels.append(f"[{label}]")
    filters.append(f"{''.join(labels)}amix=inputs={len(labels)}:duration=longest:normalize=0,atrim=duration={total_duration}[audio]")
    command += [
        "-filter_complex", ";".join(filters), "-map", "0:v:0", "-map", "[audio]",
        "-c:v", "copy", "-c:a", "aac", "-b:a", "192k", "-shortest", "-movflags", "+faststart",
        str(output_path),
    ]
    _run_ffmpeg(command)


def exportar_video(diapositivas, resolution=1080):
    output_size = (1920, 1080) if resolution == 1080 else (1280, 720)
    with tempfile.TemporaryDirectory(prefix="doit-animation-") as directory:
        directory_path = Path(directory)
        silent_video = directory_path / "silent.mp4"
        output_video = directory_path / "presentation.mp4"
        _write_animation_frames(diapositivas, silent_video, output_size)
        _mix_audio(silent_video, output_video, diapositivas, directory_path)
        return output_video.read_bytes()


def exportar_gif(diapositivas, resolution=1080):
    output_size = (1920, 1080) if resolution == 1080 else (1280, 720)
    with tempfile.TemporaryDirectory(prefix="doit-animation-") as directory:
        output_path = Path(directory) / "presentation.gif"
        _write_animation_frames(diapositivas, output_path, output_size, gif=True)
        return output_path.read_bytes()
